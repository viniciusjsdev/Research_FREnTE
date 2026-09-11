from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
REFERENCE = Path(r"C:\Users\vinic\Downloads\Do ambiente virtual à imersão internacional.docx")
OUTPUT = ROOT / "Do Rio ao Dashboard - Artigo Projeto 100K.docx"
FIGURES = ROOT / "EDA" / "operacao_reservatorio" / "figures"
PRESERVE_TEMPLATE_PARTS = {
    "word/header1.xml",
    "word/footer1.xml",
    "word/_rels/footer1.xml.rels",
    "word/styles.xml",
    "word/numbering.xml",
    "word/settings.xml",
    "word/_rels/fontTable.xml.rels",
}


def set_run_font(run, size=12, bold=None, italic=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def remove_numbering(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), 12, bold=True)
    return p


def add_heading(doc, text, numbered=True):
    p = doc.add_paragraph(style="Heading 1")
    remove_numbering(p)
    p.paragraph_format.left_indent = None
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper() if numbered else text)
    set_run_font(r, 12, bold=True)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), 12, bold=True)
    return p


def add_body(doc, text, *, indent=True, after=2, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.2958) if indent else None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run(text), size)
    return p


def add_keywords(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(label), 12, bold=True)
    set_run_font(p.add_run(text), 12)
    return p


def add_formula(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(
        p.add_run("TR (dias) = [volume útil (%) × capacidade máxima (hm³) × 10⁶] / [vazão defluente (m³/s) × 86.400]"),
        10,
        italic=True,
    )
    return p


def add_figure(doc, filename, width_in, caption, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIGURES / filename), width=Inches(width_in))
    shape._inline.docPr.set("descr", alt)

    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = False
    set_run_font(cap.add_run(caption), 9)
    return cap


def add_reference(doc, text, italic_part=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if italic_part and italic_part in text:
        before, after = text.split(italic_part, 1)
        set_run_font(p.add_run(before), 8.5)
        set_run_font(p.add_run(italic_part), 8.5, italic=True)
        set_run_font(p.add_run(after), 8.5)
    else:
        set_run_font(p.add_run(text), 8.5)
    return p


def clear_body_keep_section(doc):
    body = doc._element.body
    sectpr = body.sectPr
    for child in list(body):
        if child is not sectpr:
            body.remove(child)


def insert_cloned_paragraphs_before_section(doc, elements):
    body = doc._element.body
    sect_index = list(body).index(body.sectPr)
    for element in elements:
        body.insert(sect_index, deepcopy(element))
        sect_index += 1


def restore_template_parts():
    """Keep the reference's visual infrastructure byte-for-byte after python-docx saves."""
    temporary = OUTPUT.with_suffix(".template-preserved.docx")
    with ZipFile(OUTPUT, "r") as generated, ZipFile(REFERENCE, "r") as reference, ZipFile(
        temporary, "w", compression=ZIP_DEFLATED
    ) as rebuilt:
        for item in generated.infolist():
            payload = (
                reference.read(item.filename)
                if item.filename in PRESERVE_TEMPLATE_PARTS
                else generated.read(item.filename)
            )
            rebuilt.writestr(item, payload)
    temporary.replace(OUTPUT)


def main():
    doc = Document(REFERENCE)
    author_block = [deepcopy(doc.paragraphs[i]._p) for i in range(4, 9)]
    clear_body_keep_section(doc)

    doc.core_properties.title = (
        "Do Rio ao Dashboard: integração e análise de dados ambientais em reservatórios do Rio Tietê"
    )
    doc.core_properties.subject = "Artigo técnico do Projeto 100K"

    add_title(
        doc,
        "DO RIO AO DASHBOARD: INTEGRAÇÃO E ANÁLISE DE DADOS AMBIENTAIS EM RESERVATÓRIOS DO RIO TIETÊ NO CONTEXTO DO PROJETO 100K",
    )
    doc.add_paragraph()
    add_title(
        doc,
        "FROM THE RIVER TO THE DASHBOARD: INTEGRATION AND ANALYSIS OF ENVIRONMENTAL DATA FROM TIETÊ RIVER RESERVOIRS IN THE CONTEXT OF PROJECT 100K",
    )
    doc.add_paragraph()
    insert_cloned_paragraphs_before_section(doc, author_block)

    add_heading(doc, "RESUMO", numbered=False)
    add_body(
        doc,
        "Este trabalho apresenta uma aplicação de Ciência de Dados às Geociências desenvolvida no contexto internacional do Projeto 100K. O objetivo foi organizar, integrar e explorar dados ambientais sobre reservatórios em cascata no Rio Tietê e pressões antrópicas em seu corredor. Foram processados 56.982 registros diários do ONS para seis reservatórios (2000-2025), 1.116.805 detecções do Programa Queimadas/INPE (2000-2024), indicadores do SNIS (2010-2020) e inventários de resíduos da CETESB (2017-2021). O fluxo separou dados brutos, staging, camadas analíticas e visualizações. A vazão afluente média foi de 362 m³/s em Barra Bonita e 754 m³/s em Três Irmãos. O tempo de residência estimado apresentou medianas mensais de 15 a 204 dias e associação positiva com o volume útil (r diário de 0,16 a 0,61). A integração de fontes heterogêneas subsidia hipóteses e análises posteriores de qualidade da água, matéria orgânica e sedimentos, sem demonstrar causalidade entre pressões e respostas ambientais.",
        indent=False,
        after=4,
    )
    add_keywords(
        doc,
        "Palavras-chave: ",
        "Ciência de Dados; Geociências; reservatórios em cascata; Rio Tietê; dados ambientais.",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run("ABSTRACT"), 12, bold=True)
    add_body(
        doc,
        "This paper presents a Data Science application to Geosciences developed within Project 100K. The objective was to integrate environmental data on Tietê River cascade reservoirs and selected anthropogenic pressures. The workflow processed 56,982 daily ONS records for six reservoirs (2000-2025), 1,116,805 active-fire detections from INPE (2000-2024), SNIS sanitation indicators (2010-2020), and CETESB landfill inventories (2017-2021), keeping raw, staging, analytical, and visualization layers separate. Mean inflow was 362 m³/s at Barra Bonita and 754 m³/s at Três Irmãos. Estimated monthly median residence time ranged from 15 to 204 days and was positively associated with useful storage (daily r from 0.16 to 0.61). Integrating heterogeneous institutional data supports hypothesis generation and subsequent analyses of water quality, organic matter, and sediments, without establishing causal links between pressures and environmental responses.",
        indent=False,
        after=4,
    )
    add_keywords(
        doc,
        "Keywords: ",
        "Data Science; Geosciences; cascade reservoirs; Tietê River; environmental data.",
    )

    add_heading(doc, "1 INTRODUÇÃO")
    add_body(
        doc,
        "Sistemas fluviais regulados por barragens integram processos hidrológicos, geomorfológicos, biogeoquímicos e sociais em escalas distintas (BARBOSA et al., 1999). No Rio Tietê, reservatórios em cascata recebem sinais acumulados de urbanização, saneamento, atividades agropecuárias, resíduos e queimadas, ao mesmo tempo em que sua operação altera armazenamento e vazões. A interpretação desse sistema exige combinar séries produzidas por instituições diferentes, com formatos, recortes e graus de completude próprios.",
    )
    add_body(
        doc,
        "O Projeto 100K/FREnTE criou um contexto de internacionalização no qual estudantes de Ciência de Dados participaram de atividades de pesquisa em Geociências no Brasil e nos Estados Unidos, articulando dimensões internacional e intercultural ao ensino e à pesquisa (DE WIT et al., 2015; RESEARCH_FRENTE, 2026). Neste artigo, a mobilidade é tratada como ambiente formativo, enquanto o objeto central é a abordagem técnica desenvolvida para o eixo brasileiro. O recorte abrange o corredor do Tietê entre São Paulo e a região de Três Lagoas/Jupiá; contudo, a análise operacional implementada inclui Barra Bonita, Bariri, Ibitinga, Promissão, Nova Avanhandava e Três Irmãos, sem incorporar Jupiá na mesma rotina.",
    )
    add_body(
        doc,
        "O objetivo foi avaliar como técnicas de Ciência de Dados podem transformar registros ambientais heterogêneos em séries, indicadores e visualizações reproduzíveis. A análise não mede diretamente matéria orgânica ou sedimentos. Ela caracteriza a base hidrológica e algumas pressões que deverão ser posteriormente cruzadas com dados de qualidade da água e evidências sedimentológicas.",
    )

    add_heading(doc, "2 MATERIAIS E MÉTODOS")
    add_subheading(doc, "2.1 Arquitetura de dados e fontes")
    add_body(
        doc,
        "Foi estruturado o repositório Research_FREnTE, composto por um pipeline de descoberta de fontes, rotinas de coleta operacional e uma área de análise exploratória. A organização seguiu a sequência fonte original → raw → staging → analytic → visualização. A separação preserva a rastreabilidade do dado bruto, permite padronizar tipos e chaves antes da análise e evita que arquivos coletados sejam confundidos com produtos derivados (RESEARCH_FRENTE, 2026).",
    )
    add_body(
        doc,
        "A principal base hidrológica foi o conjunto Dados Hidráulicos por Reservatório - Base diária, do ONS, descrito pelo órgão como série de valores verificados e consolidados diariamente (ONS, 2026). Arquivos anuais de 2000 a 2025 foram filtrados pela bacia do Tietê e pelos seis reservatórios selecionados. Foram utilizados data, vazão afluente, vazão defluente e volume útil. Para pressões ambientais, foram processados focos de calor do Programa Queimadas/INPE no retângulo latitudinal -24,0 a -20,5 e longitudinal -52,2 a -45,8 (INPE, 2026); indicadores IN015 e IN046 do SNIS para prestadores locais públicos de São Paulo (BRASIL, 2022); e o Índice de Qualidade de Aterro de Resíduos (IQR) dos inventários CETESB (CETESB, 2022).",
    )
    add_subheading(doc, "2.2 Processamento e análise")
    add_body(
        doc,
        "Os registros do ONS foram convertidos para séries mensais. Calcularam-se médias mensais de volume útil e vazão afluente, média móvel de 12 meses para a vazão e mediana mensal do tempo de residência estimado. Este último foi derivado do volume útil percentual, de capacidades máximas aproximadas codificadas para cada reservatório e da vazão defluente:",
    )
    add_formula(doc)
    add_body(
        doc,
        "Valores inferiores a 1 dia ou superiores a 1.000 dias foram excluídos das visualizações de residência. A associação entre volume útil e tempo de residência foi descrita pelo coeficiente de Pearson nos registros diários válidos. As pressões ambientais foram agregadas por ano; para o SNIS e o IQR, adotou-se a mediana municipal. O processamento foi implementado em Python com pandas, NumPy, Matplotlib e Parquet, e os resultados foram publicados em apresentação HTML no GitHub Pages (RESEARCH_FRENTE, 2026).",
    )

    add_figure(
        doc,
        "fig2_vazao_cabeceira_foz.png",
        5.0,
        "Figura 1. Vazão afluente mensal e média móvel de 12 meses em Barra Bonita e Três Irmãos, 2000-2025. Fonte: elaboração própria a partir de ONS Dados Abertos.",
        "Séries de vazão afluente mensal de Barra Bonita e Três Irmãos entre 2000 e 2025.",
    )

    add_heading(doc, "3 RESULTADOS E DISCUSSÃO")
    add_subheading(doc, "3.1 Dinâmica hidrológica e tempo de residência")
    add_body(
        doc,
        "A camada de staging reuniu 56.982 registros diários entre 1º de janeiro de 2000 e 31 de dezembro de 2025, obtidos na base hidrológica do ONS (ONS, 2026). A vazão afluente média foi de 362 m³/s em Barra Bonita e 754 m³/s em Três Irmãos, aproximadamente 2,1 vezes maior no ponto mais a jusante do conjunto analisado. A Figura 1 mostra sazonalidade acentuada, picos interanuais e a redução simultânea das médias móveis no intervalo de 2014-2015. Esse comportamento descreve o regime hidrológico disponível aos reservatórios, mas não permite atribuir isoladamente suas variações à operação ou ao clima.",
    )
    add_body(
        doc,
        "As medianas mensais do tempo de residência estimado foram de 15 dias em Bariri, 19 em Ibitinga, 99 em Barra Bonita e Promissão, 108 em Nova Avanhandava e 204 dias em Três Irmãos. Esses valores não são diretamente equivalentes entre todas as usinas: as séries de volume útil de Barra Bonita, Promissão e Três Irmãos cobrem os 312 meses, enquanto Bariri, Ibitinga e Nova Avanhandava dispõem de apenas 48 meses, entre 2022 e 2025 (ONS, 2026). Em Três Irmãos, a mediana superior a 200 dias evidencia maior potencial de permanência da água no recorte, mas a estimativa depende das capacidades aproximadas e não substitui modelagem hidrodinâmica.",
    )
    add_body(
        doc,
        "A associação diária entre volume útil e tempo de residência foi positiva nos seis reservatórios, variando de r=0,16 em Bariri a r=0,61 em Ibitinga (Figura 2). Os demais coeficientes foram 0,27 em Barra Bonita, 0,34 em Promissão, 0,21 em Nova Avanhandava e 0,32 em Três Irmãos. O resultado é coerente com a própria construção de TR, que incorpora o volume no numerador, e deve ser interpretado como associação operacional, não como demonstração de causalidade ecológica. Em estudo específico de Ibitinga, Londe et al. (2016) também calcularam residência a partir de dados de vazão do ONS e mostraram que sua variação pode ser relevante para a distribuição de florações; no presente trabalho, entretanto, não foram analisadas clorofila ou florações.",
    )

    add_figure(
        doc,
        "fig4_correlacao_vol_tr.png",
        4.9,
        "Figura 2. Relação diária entre volume útil e tempo de residência estimado, com regressão linear e coeficiente de Pearson. Fonte: elaboração própria a partir de ONS Dados Abertos. Nota: Bariri, Ibitinga e Nova Avanhandava possuem volume apenas em 2022-2025.",
        "Dispersões de volume útil e tempo de residência para seis reservatórios do Tietê.",
    )

    add_subheading(doc, "3.2 Pressões ambientais e integração analítica")
    add_body(
        doc,
        "No recorte espacial do corredor, foram consolidadas 1.116.805 detecções de focos de calor entre 2000 e 2024, com pico de 182.974 em 2024, a partir do Programa Queimadas (INPE, 2026). Como a base combina satélites ativos, as contagens representam detecções, e não incêndios individualizados; mudanças na constelação de sensores limitam comparações temporais diretas. Ainda assim, a camada fornece datas e coordenadas para cruzamentos futuros com sub-bacias, períodos de baixa vazão e pontos de amostragem.",
    )
    add_body(
        doc,
        "A série do SNIS reuniu 2.429 registros municipais de 2010 a 2020. A mediana do IN015 passou de 80,0% para 83,85%, e a do IN046, de 73,76% para 80,0%; o número de municípios variou de 208 a 232 por ano (BRASIL, 2022). Esses resultados descrevem somente prestadores locais públicos de São Paulo e não constituem estimativa direta da carga lançada no Tietê. Nos inventários CETESB, o IQR mediano das UGRHIs selecionadas foi 9,2 em 2019 e 2020 e 8,9 em 2021 (CETESB, 2022). O indicador avalia a condição dos sistemas de disposição, não a presença de contaminantes na água; além disso, o recorte Tietê não foi recuperado para 2017-2018 pelo processamento atual.",
    )

    add_body(
        doc,
        "A principal contribuição do trabalho é a integração: vazão, armazenamento, residência, fogo, saneamento e resíduos passam a compartilhar uma arquitetura auditável e produtos visuais comuns. O dashboard não encerra a investigação; funciona como instrumento de exploração, comunicação e identificação de lacunas. Relações quantitativas entre essas pressões e matéria orgânica, qualidade da água ou sedimentos dependem de novas bases e desenhos analíticos específicos.",
    )

    add_heading(doc, "4 CONSIDERAÇÕES FINAIS")
    add_body(
        doc,
        "A experiência vinculada ao Projeto 100K demonstrou que a Ciência de Dados pode atuar como ponte entre diferentes fontes ambientais e a interpretação geocientífica de reservatórios em cascata. A arquitetura implementada transformou arquivos institucionais heterogêneos em camadas rastreáveis, séries temporais, indicadores e visualizações acessíveis, além de explicitar diferenças de escala e de cobertura.",
    )
    add_body(
        doc,
        "Os resultados caracterizaram a ampliação da vazão ao longo da cascata, a heterogeneidade do tempo de residência estimado e pressões ambientais selecionadas. As limitações encontradas - cobertura desigual de volume útil, caráter derivado da residência, diferenças amostrais do SNIS e comparabilidade dos sensores de fogo - são parte do resultado técnico e orientam a próxima etapa. O avanço científico requer integrar esta base a séries de qualidade da água, matéria orgânica e sedimentos, preservando a distinção entre associação, hipótese e causalidade.",
    )

    add_heading(doc, "REFERÊNCIAS", numbered=False)
    add_reference(
        doc,
        "BARBOSA, F. A. R.; PADISÁK, J.; ESPÍNDOLA, E. L. G.; BORICS, G.; ROCHA, O. The cascading reservoir continuum concept (CRCC) and its application to the River Tietê-basin, São Paulo State, Brazil. In: TUNDISI, J. G.; STRAŠKRABA, M. (ed.). Theoretical Reservoir Ecology and its Applications. São Carlos: Brazilian Academy of Sciences; International Institute of Ecology; Backhuys Publishers, 1999. p. 425-437.",
        italic_part="Theoretical Reservoir Ecology and its Applications",
    )
    add_reference(
        doc,
        "BRASIL. Ministério das Cidades. Sistema Nacional de Informações sobre Saneamento. Glossário de Indicadores - Água e Esgotos 2021. Brasília, DF, 2022. Disponível em: https://www.gov.br/cidades/pt-br/acesso-a-informacao/acoes-e-programas/saneamento/snis/produtos-do-snis/diagnosticos/Glossario_Indicadores_AE2021.pdf. Acesso em: 2 set. 2026.",
    )
    add_reference(
        doc,
        "CETESB - Companhia Ambiental do Estado de São Paulo. Inventário estadual de resíduos sólidos urbanos 2021. São Paulo: CETESB, 2022. Disponível em: https://repositorio.cetesb.sp.gov.br/handle/123456789/2493. Acesso em: 2 set. 2026.",
        italic_part="Inventário estadual de resíduos sólidos urbanos 2021",
    )
    add_reference(
        doc,
        "DE WIT, H.; HUNTER, F.; HOWARD, L.; EGRON-POLAK, E. Internationalisation of higher education. Brussels: European Parliament, 2015. DOI: 10.2861/444393.",
        italic_part="Internationalisation of higher education",
    )
    add_reference(
        doc,
        "INPE - Instituto Nacional de Pesquisas Espaciais. Programa Queimadas: dados para download. São José dos Campos, 2026. Disponível em: https://terrabrasilis.dpi.inpe.br/queimadas/portal/. Acesso em: 2 set. 2026.",
    )
    add_reference(
        doc,
        "LONDE, L. R.; NOVO, E. M. L. M.; BARBOSA, C.; ARAUJO, C. A. S. Water residence time affecting phytoplankton blooms: study case in Ibitinga Reservoir (São Paulo, Brazil) using Landsat/TM images. Brazilian Journal of Biology, v. 76, n. 3, p. 664-672, 2016. DOI: 10.1590/1519-6984.23814.",
        italic_part="Brazilian Journal of Biology",
    )
    add_reference(
        doc,
        "ONS - Operador Nacional do Sistema Elétrico. Dados Hidráulicos por Reservatório - Base diária. Rio de Janeiro, 2026. Disponível em: https://dados.ons.org.br/dataset/dados-hidrologicos-res. Acesso em: 2 set. 2026.",
    )
    add_reference(
        doc,
        "RESEARCH_FRENTE. Pipeline multiagente e análises exploratórias do Projeto 100K. GitHub, 2026. Disponível em: https://github.com/viniciusjsdev/Research_FREnTE. Dashboard: https://viniciusjsdev.github.io/Research_FREnTE/. Acesso em: 2 set. 2026.",
    )

    doc.save(OUTPUT)
    restore_template_parts()
    print(OUTPUT)


if __name__ == "__main__":
    main()
