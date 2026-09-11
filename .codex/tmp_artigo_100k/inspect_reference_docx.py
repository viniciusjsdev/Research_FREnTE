from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


REF = Path(r"C:\Users\vinic\Downloads\Do ambiente virtual à imersão internacional.docx")
OUT = Path(__file__).resolve().parent


def emu_to_in(value):
    return round(value / 914400, 4) if value is not None else None


def font_name(run):
    if run.font.name:
        return run.font.name
    rpr = run._element.rPr
    if rpr is not None and rpr.rFonts is not None:
        return rpr.rFonts.get(qn("w:ascii")) or rpr.rFonts.get(qn("w:hAnsi"))
    return None


def main() -> None:
    doc = Document(REF)
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        pf = p.paragraph_format
        runs = []
        for r in p.runs:
            runs.append(
                {
                    "text": r.text,
                    "font": font_name(r),
                    "size_pt": round(r.font.size.pt, 2) if r.font.size else None,
                    "bold": r.bold,
                    "italic": r.italic,
                    "superscript": r.font.superscript,
                }
            )
        paragraphs.append(
            {
                "index": i,
                "style": p.style.name if p.style else None,
                "text": p.text,
                "alignment": str(p.alignment),
                "left_in": emu_to_in(pf.left_indent),
                "right_in": emu_to_in(pf.right_indent),
                "first_line_in": emu_to_in(pf.first_line_indent),
                "space_before_pt": round(pf.space_before.pt, 2) if pf.space_before else None,
                "space_after_pt": round(pf.space_after.pt, 2) if pf.space_after else None,
                "line_spacing": str(pf.line_spacing),
                "keep_with_next": pf.keep_with_next,
                "page_break_before": pf.page_break_before,
                "runs": runs,
            }
        )

    sections = []
    for s in doc.sections:
        sections.append(
            {
                "page_width_in": emu_to_in(s.page_width),
                "page_height_in": emu_to_in(s.page_height),
                "top_margin_in": emu_to_in(s.top_margin),
                "bottom_margin_in": emu_to_in(s.bottom_margin),
                "left_margin_in": emu_to_in(s.left_margin),
                "right_margin_in": emu_to_in(s.right_margin),
                "header_distance_in": emu_to_in(s.header_distance),
                "footer_distance_in": emu_to_in(s.footer_distance),
                "different_first_page": s.different_first_page_header_footer,
            }
        )

    with zipfile.ZipFile(REF) as zf:
        parts = []
        for info in zf.infolist():
            parts.append(
                {
                    "path": info.filename,
                    "size": info.file_size,
                    "sha256": hashlib.sha256(zf.read(info.filename)).hexdigest(),
                }
            )

    payload = {
        "reference": str(REF),
        "sha256": hashlib.sha256(REF.read_bytes()).hexdigest(),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "sections": sections,
        "paragraphs": paragraphs,
        "parts": parts,
    }
    (OUT / "reference_inventory.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({k: payload[k] for k in ("reference", "sha256", "paragraph_count", "table_count", "inline_shapes", "sections")}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
